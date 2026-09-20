from pathlib import Path
from typing import Dict, List, Sequence, Union
import openpyxl  # type: ignore
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
from openpyxl.utils.dataframe import dataframe_to_rows  # type: ignore
import pandas as pd

from sipangpt_eval.config import settings
from sipangpt_eval.schemas.evaluation import BenchmarkSummaryMetrics, EvaluatedPair


def generate_comparative_excel(
    pairs: Sequence[EvaluatedPair],
    metrics: BenchmarkSummaryMetrics,
    output_path: Path = settings.results_dir / "matriz_comparativa_final.xlsx",
) -> None:
    """Genera la matriz comparativa final en formato Excel (.xlsx) con diseño profesional."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows_detail: List[Dict[str, Union[int, float, str]]] = []

    for pair in pairs:
        c = pair.caso
        ft = pair.inferencia_finetuned
        rag = pair.inferencia_rag

        citas_str: str = "; ".join(
            [f"{cit.documento} ({cit.articulo_o_seccion or ''})" for cit in rag.citas]
        )

        rows_detail.append(
            {
                "ID": c.id,
                "Modulo": c.modulo.value,
                "Tipo": c.tipo.value.capitalize(),
                "Pregunta_Usuario": c.pregunta_usuario,
                "Respuesta_Esperada_GroundTruth": c.respuesta_esperada,
                "Respuesta_Gemma4_Finetuned": ft.respuesta_generada,
                "Latencia_FT_ms": round(ft.tiempo_total_ms, 2),
                "Calificacion_Finetuned": pair.score_finetuned.value,
                "Respuesta_Sipan_STAIR_RAG": rag.respuesta_generada,
                "Latencia_RAG_ms": round(rag.tiempo_total_ms, 2),
                "Citas_RAG": citas_str,
                "Calificacion_RAG": pair.score_rag.value,
                "Observaciones": pair.observaciones,
            }
        )

    df_detail = pd.DataFrame(rows_detail)

    rows_metrics: List[Dict[str, Union[int, float, str]]] = [
        {"Métrica": "Total Preguntas de Prueba", "Valor": metrics.total_preguntas},
        {
            "Métrica": "Exactitud Gemma-4 Fine-Tuned (%)",
            "Valor": f"{metrics.exactitud_finetuned_pct:.2f}%",
        },
        {
            "Métrica": "Exactitud Sipán-STAIR RAG (%)",
            "Valor": f"{metrics.exactitud_rag_pct:.2f}%",
        },
        {
            "Métrica": "Total Alucinaciones (Gemma-4 Fine-Tuned)",
            "Valor": metrics.alucinaciones_finetuned_count,
        },
        {
            "Métrica": "Total Alucinaciones (Sipán-STAIR RAG)",
            "Valor": metrics.alucinaciones_rag_count,
        },
        {
            "Métrica": "Latencia Media Gemma-4 Fine-Tuned (ms)",
            "Valor": metrics.latencia_media_finetuned_ms,
        },
        {
            "Métrica": "Latencia Media Sipán-STAIR RAG (ms)",
            "Valor": metrics.latencia_media_rag_ms,
        },
    ]

    df_metrics = pd.DataFrame(rows_metrics)

    wb = openpyxl.Workbook()
    ws_detail = wb.active
    ws_detail.title = "Matriz_Comparativa_50_Casos"

    for r in dataframe_to_rows(df_detail, index=False, header=True):
        ws_detail.append(r)

    ws_summary = wb.create_sheet(title="Resumen_Metricas")
    for r in dataframe_to_rows(df_metrics, index=False, header=True):
        ws_summary.append(r)

    # Aplicar estilos visuales profesionales
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for sheet in [ws_detail, ws_summary]:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in sheet.columns:
            max_len: int = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(output_path)
